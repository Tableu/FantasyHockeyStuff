"""Steve Laidlaw Fantasy Hockey Rankings -- skaters only, no goalie projections.
The source's own Notes sheet explains its "punt goalies" strategy: no goalie stat
projections are provided at all, only tier labels (Goalie Tiers sheet, not read here).
No team column and no FOW/FOL/+/-/PPG/PPA/PIM/TOI/SHP columns exist either, so those
stay unset. Uses the Skaters sheet (union of Forwards+Defense) rather than reading the
two split tabs separately.
"""

from pathlib import Path

import openpyxl

FILENAME = "Steve Laidlaw Fantasy Hockey Rankings.xlsx"
SHEET = "Skaters"
HEADER_ROW = 1
FIRST_DATA_ROW = 2


def _to_int(value):
    return int(round(float(value))) if value not in (None, "") else None


def _header_map(ws, row):
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=row, column=c).value
        if isinstance(v, str) and v.strip() and v.strip() not in out:
            out[v.strip()] = c
    return out


def rows(sheets_dir: Path):
    wb = openpyxl.load_workbook(sheets_dir / FILENAME, data_only=True)
    ws = wb[SHEET]
    hmap = _header_map(ws, HEADER_ROW)

    def get(row, header):
        return row[hmap[header] - 1]

    for r in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        name = r[0]
        if not name:
            continue
        yield {
            "raw_name": name,
            "team_raw": None,
            "is_goalie": False,
            "stats": {
                "GamesPlayed": _to_int(get(r, "GP")),
                "Goals": _to_int(get(r, "G")),
                "Assists": _to_int(get(r, "A")),
                "Points": _to_int(get(r, "P")),
                "PowerPlayPoints": _to_int(get(r, "PPP")),
                "Shots": _to_int(get(r, "SOG")),
                "Hits": _to_int(get(r, "Hits")),
                "Blocks": _to_int(get(r, "Blks")),
            },
        }
